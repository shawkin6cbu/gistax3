import os
from tkinter import Tk
from docx import Document
from gistax.desoto.data import SharedData
from gistax.desoto.gui.processing_tab import ProcessingTab
from gistax.desoto.services.title_chain import extract_table_entries_from_pdf, get_24_month_chain
from datetime import datetime

def main():
    root = Tk()
    root.withdraw()
    shared = SharedData()
    tab = ProcessingTab(root, shared)

    # Mock data
    tab.pin_var.set("123-45-6789.000")
    tab.address_var.set("123 Mockingbird Ln")
    tab.owner_var.set("JOHN DOE & JANE DOE")
    tab.city_var.set("Hernando, MS 38632")
    tab.legal_desc_var.set("Lot 12, SOME SUBDIVISION")

    tab.lender_var.set("ACME BANK")
    tab.borrower_var.set("JANE DOE")

    tab.tax_2024_total_var.set("1,234.56")
    tab.tax_2024_date_paid_var.set("03/15/2025")
    tab.tax_2025_est_var.set("1,300.00")

    # Build title chain from test PDF
    pdf_path = os.path.join(os.path.dirname(__file__), 'test', 'chain_test_page.pdf')
    entries = extract_table_entries_from_pdf(pdf_path)
    chain = get_24_month_chain(entries, datetime(2025, 8, 10))
    shared.set_data('title_chain_results', chain)

    # Generate document
    out_path = os.path.abspath(os.path.join(os.path.dirname(__file__), 'TitleDocs.docx'))
    ok, msg = tab._create_full_document(out_path)
    print('GENERATE_OK', ok)
    print(msg)

    # Verify placeholders were replaced
    doc = Document(out_path)
    full_text = '\n'.join(p.text for p in doc.paragraphs)
    for token in ["{PARCEL}", "{PROPSTRE}", "{SLRLAST}", "{TAXAMT}", "{TAXDAT}", "{BYRLAST}", "{Lender}"]:
        if token in full_text:
            print('TOKEN_STILL_PRESENT', token)

    # Verify some fields present
    checks = [
        "123-45-6789.000",
        "123 Mockingbird Ln",
        "JOHN DOE & JANE DOE",
        "Hernando, MS 38632",
        "1,234.56",
        "03/15/2025",
        "ACME BANK",
        "JANE DOE",
    ]
    for c in checks:
        print('HAS', c, c in full_text)

if __name__ == '__main__':
    main()

